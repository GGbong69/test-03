#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# 채운 이름표를 data/items.csv 의 name 열에 되돌린다.
#
#     python scripts/tools/apply_names.py --dry     무엇이 바뀌는지만 본다
#     python scripts/tools/apply_names.py           실제로 쓴다
#     python scripts/tools/apply_names.py 다른표.xlsx
#
# 빈칸은 건드리지 않는다 — 150장을 한 번에 안 채워도 되고, 채운 만큼만
# 여러 번 돌려도 된다.
#
# name 열 하나만 바꾼다. 줄을 앞에서 두 번만 쪼개고(id, name, 나머지)
# 나머지는 손대지 않으므로 다른 열의 서식이 통째로 보존된다.
#
# 검사는 첫 오류에서 안 멈춘다. 전부 모아 한 번에 내놓고, 오류가 하나라도
# 있으면 아무것도 안 쓴다 — 반만 들어간 표가 제일 고약하다.
import argparse
import os
import sys

import openpyxl

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except AttributeError:
        pass

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CSV = os.path.join(ROOT, "data", "items.csv")
XLSX = os.path.join(ROOT, "shots", "표", "이름.xlsx")

H_ID = "id"
H_OLD = "지금 이름"
H_NEW = "새 이름"
HEAD_MAX = 6        # 머리글을 이 줄 안에서 찾는다


def find_head(ws):
    """머리글 줄과 열의 자리를 찾는다. 사람이 위에 줄을 끼워 넣어도 깨지지
    않게 이름으로 찾는다 — 자리를 외우면 표가 조금만 움직여도 죽는다.

    id 열은 표에서 숨겨져 있다. 사람이 볼 것이 아니라 줄을 짚는 열쇠라서
    보일 이유가 없다. 다른 시트로 옮겨 붙이다 id 가 통째로 빠졌으면
    '지금 이름' 으로도 짚는다 — 이름은 185장이 서로 다르다."""
    for r in range(1, min(HEAD_MAX, ws.max_row) + 1):
        head = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                head[v.strip()] = c
        if H_NEW in head and (H_ID in head or H_OLD in head):
            return r, head.get(H_ID), head.get(H_OLD), head[H_NEW]
    return None, None, None, None


def read_sheet(ws, by_name, errs):
    """(id, 새 이름, 어디) 목록. 빈칸은 안 담는다."""
    hr, ci, co, cn = find_head(ws)
    if hr is None:
        errs.append("시트 '%s' — 머리글에 '%s' 와 ('%s' 또는 '%s') 가 같이 "
                    "있는 줄이 없다" % (ws.title, H_NEW, H_ID, H_OLD))
        return []
    got = []
    for r in range(hr + 1, ws.max_row + 1):
        new = ws.cell(r, cn).value
        if new is None or str(new).strip() == "":
            continue
        where = "%s!%d행" % (ws.title, r)
        iid = str(ws.cell(r, ci).value or "").strip() if ci else ""
        if not iid and co:
            # id 가 빠졌으면 지금 이름으로 짚는다
            old = str(ws.cell(r, co).value or "").strip()
            iid = by_name.get(old, "")
            if not iid:
                errs.append("%s — id 도 없고 '지금 이름' 도 못 찾겠다: %r"
                            % (where, old))
                continue
        if not iid:
            errs.append("%s — 줄을 짚을 id 가 없다" % where)
            continue
        got.append((iid, str(new).strip(), where))
    return got


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=XLSX)
    ap.add_argument("--dry", action="store_true", help="쓰지 않고 보기만")
    ap.add_argument("--max", type=int, default=7,
                    help="이 글자수를 넘으면 경고(기본 7)")
    a = ap.parse_args()

    if not os.path.exists(a.xlsx):
        sys.exit("없다: %s — 먼저 make_name_xlsx.py 로 뽑아라" % a.xlsx)

    raw = open(CSV, encoding="utf-8").read().split("\n")
    order = []          # 파일에 있는 차례대로의 id
    cur = {}            # id → 지금 이름
    at = {}             # id → 몇 번째 줄
    for i, line in enumerate(raw):
        if i == 0 or line.strip() == "":
            continue
        p = line.split(",", 2)
        if len(p) < 3:
            continue
        order.append(p[0])
        cur[p[0]] = p[1]
        at[p[0]] = i

    by_name = {}
    for iid, nm in cur.items():
        by_name[nm] = iid

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    errs, warns, want, seen = [], [], {}, {}
    for ws in wb.worksheets:
        for iid, new, where in read_sheet(ws, by_name, errs):
            if iid not in cur:
                errs.append("%s — 모르는 id '%s'" % (where, iid))
                continue
            if iid in seen:
                errs.append("%s — id '%s' 가 %s 에도 있다" % (where, iid, seen[iid]))
                continue
            seen[iid] = where
            if "," in new:
                errs.append("%s — '%s' 에 쉼표가 있다. 표가 깨진다" % (where, new))
                continue
            if new == cur[iid]:
                continue                    # 그대로다 — 바꿀 것이 없다
            want[iid] = new
            if len(new) > a.max:
                warns.append("%s — '%s' 가 %d글자다. 정산 카드가 이름과 효과를 "
                             "한 줄에 그린다" % (where, new, len(new)))

    # 다 바꾼 뒤의 이름으로 중복을 본다. 안 바꾼 이름과 부딪히는 것도 잡아야
    # 한다 — 표 안에서만 보면 기존 35장과의 충돌을 놓친다.
    final = {}
    for iid in order:
        nm = want.get(iid, cur[iid])
        final.setdefault(nm, []).append(iid)
    for nm, ids in final.items():
        if len(ids) > 1:
            errs.append("이름 '%s' 가 겹친다: %s" % (nm, " · ".join(ids)))

    for w in warns:
        print("경고: " + w)
    for e in errs:
        print("오류: " + e)
    if errs:
        sys.exit("\n오류 %d건 — 아무것도 안 썼다." % len(errs))

    if not want:
        print("바꿀 것이 없다. 채운 칸이 %d개." % len(seen))
        return

    for iid in order:
        if iid in want:
            print("  %-6s %s → %s" % (iid, cur[iid], want[iid]))
    print("%d장 (채운 칸 %d개 중 %d개는 지금과 같다)"
          % (len(want), len(seen), len(seen) - len(want)))

    if a.dry:
        print("--dry 라 안 썼다.")
        return

    for iid, new in want.items():
        i = at[iid]
        p = raw[i].split(",", 2)
        raw[i] = ",".join([p[0], new, p[2]])
    open(CSV, "w", encoding="utf-8", newline="").write("\n".join(raw))
    print("씀: %s" % CSV)
    print("파생 문서는 따로 다시 뽑는다 — export_kb.py · make_deck.py")


if __name__ == "__main__":
    main()
