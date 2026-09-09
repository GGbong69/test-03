#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# 동전 이름을 손으로 다시 짓기 위한 엑셀을 낸다.
#
#     python scripts/tools/make_name_xlsx.py
#
# 등급마다 시트를 하나씩 두고, 시트 안에서는 **조건끼리 모은다**. 같은
# 조건의 동전이 붙어 있어야 "이 식구들" 을 한 묶음으로 보고 계열이 서는
# 이름을 지을 수 있다. 미구현 장을 아래로 몰지 않는다 — 이름도 능력도
# 그 장들이 제일 손이 많이 가는데 밑으로 내리면 안 보인다.
#
# 능력은 한 문장으로 적는다. 낱말은 하나도 안 지어낸다 — dump_items.gd 가
# 게임의 cond_text / eff_line / gold_text 로 찍어 둔 세 도막을 잇기만 한다.
# 게임은 툴팁에서 이 셋을 세 줄로 보여주는데(game.gd:8801), 이름을 지을
# 때는 한 줄로 읽히는 편이 낫다.
#
# 86장은 능력 칸이 빈다. 없어서가 아니라 **엔진이 아직 못 해서** 눕혀 둔
# 장들이고(enabled=0), 원본 효과 문구는 저장소에 없다 — 이식 커밋
# 88824db 가 외부 시트에서 이름과 막힌 갈래만 받아 왔다. 그 자리를 위해
# "새 능력" 칸을 둔다.
#
# 이미 채운 것은 그대로 이어받는다. 표를 다시 뽑는다고 사람이 적어 둔
# 것이 날아가면 도구가 아니라 함정이다.
#
# 덤프가 낡았으면 먼저 다시 뽑는다:
#     godot --headless --script scripts/tools/dump_items.gd
#
# id 는 열로 두되 **숨긴다**. 사람이 볼 것은 아니지만 apply_names.py 가
# 줄을 짚는 열쇠다. 숨긴 열은 줄을 정렬해도 같이 따라간다.
import csv
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except AttributeError:
        pass

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "data", "items.csv")
DUMP = os.path.join(ROOT, "shots", "items_all.csv")
ORIG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "joker_origin.csv")
OUT = os.path.join(ROOT, "shots", "표", "이름.xlsx")

HEAD_BG = "FF217346"
FILL_HD = "FFBF8F00"
FILL_BG = "FFFFF2CC"     # 채울 칸 — 노란 자리가 곧 "여기 적어라"
DONE_BG = "FFEDEDED"     # 이미 손으로 지은 이름
THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F9 = Font(name="맑은 고딕", size=9)
F10 = Font(name="맑은 고딕", size=10)
F10G = Font(name="맑은 고딕", size=10, color="FF808080")
F10R = Font(name="맑은 고딕", size=10, color="FFB25000")

HEAD_ROW = 2
DATA_ROW = 3

RARS = [("common", "일반"), ("uncommon", "희귀"), ("rare", "레어"),
        ("legendary", "레전더리")]

H_ID = "id"
H_OLD = "지금 이름"
H_NEW = "새 이름"
H_ABIL = "새 능력"
H_ORIG = "원본 (발라트로 · 추정)"
H_BLOCK = "막힌 갈래 · 메모"

# (머리글, 너비, 왼쪽정렬, 채우는 칸인가)
COLS = [
    (H_ID, 9, False, False),        # 숨긴다 — apply_names.py 의 열쇠
    (H_OLD, 20, True, False),
    (H_NEW, 20, True, True),
    ("상태", 9, False, False),
    ("능력", 52, True, False),
    (H_ABIL, 40, True, True),
    (H_ORIG, 52, True, False),
    ("태그", 16, True, False),
    (H_BLOCK, 30, True, False),
]
ID_COL = 1


def hand_made(iid):
    """이미 손으로 지은 이름인가. 이식본은 j + 숫자다."""
    return not (iid.startswith("j") and iid[1:].isdigit())


def ability(r, d):
    """능력 한 문장. 게임이 찍은 세 도막을 잇는다.

    조건·효과가 통째로 빈 장은 문장이 안 나온다. 지어내지 않고 비워 둔다 —
    막힌 갈래는 옆 칸에 따로 적힌다."""
    cond = d.get("조건", "").strip()
    eff = d.get("효과", "").strip()
    gold = d.get("골드", "").strip()

    if not eff and not gold:
        return ""
    if not eff:
        # 골드 전업. 골드 문장이 제 조건을 이미 품고 있어서 "모든 다트" 를
        # 앞에 붙이면 조건이 둘로 읽힌다.
        return gold if cond in ("", "모든 다트") else "%s · %s" % (cond, gold)
    # 조건과 효과는 줄표로 가른다. 그냥 붙이면 "모든 다트 남은 다트 1개당"
    # 처럼 말이 엉킨다.
    s = "%s — %s" % (cond, eff) if cond else eff
    return s + (" · " + gold if gold else "")


def blocked(r):
    """엔진이 못 하는 갈래. enabled=0 인 장의 메모가 그것을 적고 있다."""
    n = r["_note"].strip()
    return n[len("엔진 미지원: "):] if n.startswith("엔진 미지원: ") else n


def origins(path):
    """이식 전 발라트로 조커의 능력. **추정이다.**

    원본 문구는 저장소에 없다 — 이식 커밋 88824db 가 외부 시트에서 이름과
    막힌 갈래만 받아 왔다. 그래서 이름을 보고 어느 조커인지 되짚은 것이고,
    확신 칸이 그 정도를 말한다. 못 짚은 장은 물음표로 남겨 두었다 —
    지어내느니 비는 편이 낫다."""
    got = {}
    if not os.path.exists(path):
        return got
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        src, eff, conf = r["원본"].strip(), r["효과"].strip(), r["확신"].strip()
        got[r["id"].strip()] = eff if src == "?" else "%s — %s%s" % (
            src, eff, "" if conf == "높음" else "  (확신 %s)" % conf)
    return got


def carry(path):
    """전에 채워 둔 새 이름 · 새 능력을 id 로 이어받는다."""
    got = {}
    if not os.path.exists(path):
        return got
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:            # 열다 깨져도 표 뽑는 것은 계속된다
        print("경고: 옛 표를 못 읽었다(%s) — 채운 것을 못 이어받는다" % e)
        return got
    for ws in wb.worksheets:
        head = {}
        for r in range(1, min(6, ws.max_row) + 1):
            head = {str(ws.cell(r, c).value).strip(): c
                    for c in range(1, ws.max_column + 1)
                    if isinstance(ws.cell(r, c).value, str)}
            if H_ID in head and H_NEW in head:
                break
        if H_ID not in head or H_NEW not in head:
            continue
        for r in range(1, ws.max_row + 1):
            iid = ws.cell(r, head[H_ID]).value
            if not iid or str(iid).strip() in ("", H_ID):
                continue
            nm = ws.cell(r, head[H_NEW]).value
            ab = ws.cell(r, head[H_ABIL]).value if H_ABIL in head else None
            if nm or ab:
                got[str(iid).strip()] = (
                    str(nm).strip() if nm else None,
                    str(ab).strip() if ab else None)
    return got


def head_cell(ws, row, col, text, bg=HEAD_BG):
    c = ws.cell(row, col, text)
    c.font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal="center")
    c.border = BOX
    return c


def build(ws, rows, dump, old, orig):
    t = ws.cell(1, 1, "노란 칸(새 이름 · 새 능력)만 채운다. 빈칸은 안 바뀐다 — "
                      "나눠서 채워도 되고, 다시 뽑아도 채운 것은 남는다. "
                      "회색 줄은 이미 손으로 지은 이름. "
                      "상태가 미구현인 장은 엔진이 아직 못 해서 눕혀 둔 것이라 "
                      "능력 칸이 비어 있다 — 오른쪽 막힌 갈래가 그 사유다. "
                      "쉼표는 못 쓴다. 되돌리기: python scripts/tools/apply_names.py")
    t.font = F9

    for j, (nm, w, _l, fill) in enumerate(COLS, start=1):
        head_cell(ws, HEAD_ROW, j, nm, FILL_HD if fill else HEAD_BG)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.column_dimensions[get_column_letter(ID_COL)].hidden = True

    for i, r in enumerate(rows):
        rr = DATA_ROW + i
        off = r["enabled"].strip() == "0"
        done = hand_made(r["id"])
        pre_n, pre_a = old.get(r["id"], (None, None))
        vals = [r["id"], r["name"].strip(), pre_n,
                "미구현" if off else "씀",
                ability(r, dump.get(r["id"], {})), pre_a,
                orig.get(r["id"], ""),
                r["_tags"].strip().replace(";", " · "),
                blocked(r) if off else r["_note"].strip()]
        for j, ((nm, _w, left, fill), v) in enumerate(zip(COLS, vals), start=1):
            c = ws.cell(rr, j, v if v not in ("", None) else None)
            c.border = BOX
            c.alignment = Alignment(horizontal="left" if left else "center",
                                    vertical="center",
                                    wrap_text=(nm in ("능력", H_ORIG)))
            if nm == H_ORIG:
                c.font = F10G
            elif nm == H_BLOCK:
                c.font = F10R if off else F10G
            elif nm == "상태" and off:
                c.font = F10R
            else:
                c.font = F10
            if fill:
                c.fill = PatternFill("solid", start_color=FILL_BG)
            elif done:
                c.fill = PatternFill("solid", start_color=DONE_BG)

    ws.freeze_panes = ws.cell(DATA_ROW, 4)
    return len(rows)


def main():
    rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig")))
    dump = {}
    if os.path.exists(DUMP):
        dump = {r["id"]: r for r in csv.DictReader(open(DUMP, encoding="utf-8-sig"))}
    miss = [r["id"] for r in rows if r["id"] not in dump]
    if miss:
        print("경고: 덤프에 없는 동전 %d장 — 능력 칸이 빈다." % len(miss))
        print("      godot --headless --script scripts/tools/dump_items.gd")

    old = carry(OUT)
    orig = origins(ORIG)

    wb = openpyxl.Workbook()
    first, total, off = True, 0, 0
    for key, title in RARS:
        got = [r for r in rows if r["rarity"].strip() == key]
        if not got:
            continue
        # 조건끼리 모으고 그 안은 효과 순. 미구현이라고 아래로 안 민다.
        got.sort(key=lambda r: (r["cond"].strip(), r["kind"].strip(),
                                int(r["value"] or 0), r["id"]))
        ws = wb.active if first else wb.create_sheet()
        ws.title = "%s(%d)" % (title, len(got))
        first = False
        total += build(ws, got, dump, old, orig)
        off += sum(1 for r in got if r["enabled"].strip() == "0")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    left = sum(1 for r in rows if not hand_made(r["id"]))
    print("씀: %s" % OUT)
    print("  %d행 · 시트 %d장 · 이름 지을 것 %d장(나머지 %d장은 이미 지었다)"
          % (total, len(wb.sheetnames), left, total - left))
    print("  미구현 %d장 — 능력 칸이 비고 막힌 갈래에 사유가 있다" % off)
    if orig:
        print("  원본(추정) %d장을 붙였다 — scripts/tools/joker_origin.csv" % len(orig))
    if old:
        print("  전에 채운 %d줄을 이어받았다" % len(old))


if __name__ == "__main__":
    main()
